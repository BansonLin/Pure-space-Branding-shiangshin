#!/usr/bin/env python3
"""Editable data-request tracker (xlsx) that mirrors pages 12–13 of the proposal PDF."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
FACTS = json.loads((ROOT / "data" / "extracted_facts.json").read_text(encoding="utf-8"))
OUT = ROOT / "reports" / "Shuangfu768_Data_Request_Tracker_v1.0.xlsx"

DUE = {"P1": "2026-11-13", "P2": "2026-11-27", "P3": "2027-01-08"}
HEAD_FILL = PatternFill("solid", fgColor="2F6B5B")
PRI_FILL = {"P1": PatternFill("solid", fgColor="F8E5E2"), "P2": PatternFill("solid", fgColor="FBF0DC"),
            "P3": PatternFill("solid", fgColor="ECEFF0")}
THIN = Side(style="thin", color="D8DDD9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = "Noto Sans CJK TC"


def style_sheet(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for cell in ws[1]:
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER


wb = Workbook()
ws = wb.active
ws.title = "資料需求追蹤"
ws.append(["#", "優先", "期限", "資料項目", "用途", "格式", "建商承接人", "狀態", "收件日", "璞域檢核人", "備註"])
for r in FACTS["data_requests"]:
    ws.append([r["no"], r["priority"], DUE[r["priority"]], r["item"], r["purpose"], r["format"], r["owner"],
               "未提供", None, "璞域專案經理", None])
style_sheet(ws, [5, 7, 12, 46, 26, 14, 20, 10, 12, 14, 24])
for row in ws.iter_rows(min_row=2):
    row[1].fill = PRI_FILL[row[1].value]
dv = DataValidation(type="list", formula1='"未提供,部分提供,已提供,不適用"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(f"H2:H{ws.max_row}")
n = ws.max_row
ws.append([])
ws.append(["", "", "", "P1 到位率（CP0 門檻 80%）", f'=COUNTIFS(B2:B{n},"P1",H2:H{n},"已提供")/COUNTIF(B2:B{n},"P1")'])
ws.cell(row=ws.max_row, column=5).number_format = "0%"
ws.cell(row=ws.max_row, column=4).font = Font(name=FONT, bold=True)

cp = wb.create_sheet("檢查點")
cp.append(["檢查點", "預定時間", "通過條件", "未通過的動作", "承接人", "結果", "日期", "備註"])
for row in [
    ["CP0 資料到位", "2026-11-15", "P1 資料到位 ≥ 80%", "暫停 Phase 1 以外工作，重排時程", "璞域專案經理"],
    ["CP1 定位定價共識", "2026-12-13", "客群、產品組合、價格區間書面確認", "視覺製作延後開工", "建商負責人＋璞域企劃總監"],
    ["G1 法規閘門", "2027-01-22", "建照核發＋預售契約與建案資訊備查（委託代銷另需代銷備查）", "不得刊登任何廣告、不收訂", "建商法務／地政士"],
    ["CP3 潛銷檢查", "潛銷第 4 週", "有效來客 ≥ 100 組、意向登記 ≥ 25 組（示意）", "調整價格帶、付款條件或主打戶型", "璞域企劃總監"],
    ["CP4 公開檢查", "公開第 8 週", "成交 ≥ 13 戶，30%（示意）", "啟動價格與通路檢討", "建商負責人＋璞域企劃總監"],
]:
    cp.append(row)
style_sheet(cp, [16, 13, 40, 30, 24, 10, 12, 24])

sc = wb.create_sheet("二工盤點")
sc.append(["位置", "請照圖用途", "二工家配圖標示", "風險（評估）", "建議做法", "建築師確認", "處理結果"])
for r in FACTS["secondary_construction"]:
    sc.append([r["where"], r["permit"], r["plan"], r["level"], r["advice"], None, None])
style_sheet(sc, [14, 26, 30, 10, 46, 14, 24])

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print("wrote", OUT)
